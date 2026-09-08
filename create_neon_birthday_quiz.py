from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTPUT = Path(__file__).with_name("Svoya_igra_Vlad_Alla.pptx")

W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(13, 10, 40)
DEEP_PURPLE = RGBColor(31, 19, 76)
CYAN = RGBColor(45, 245, 255)
PINK = RGBColor(255, 50, 184)
YELLOW = RGBColor(255, 225, 74)
WHITE = RGBColor(247, 246, 255)
MUTED = RGBColor(182, 175, 218)

QUESTIONS = [
    ("ИМЕНИННИКИ", 100, "Сколько лет сегодня каждому\nиз главных героев вечера?", "30", "Точный ответ — 30. Аплодисменты обязательны!"),
    ("ИМЕНИННИКИ", 200, "Назовите имена дуэта,\nради которого мы сегодня собрались.", "Влад и Алла", "Влад и Алла — звёзды этого вечера."),
    ("ИМЕНИННИКИ", 300, "Придумайте тост для Влада и Аллы,\nне используя слова «счастье» и «здоровье».", "Творческий раунд", "Любой тост, который рассмешит или растрогает именинников, приносит баллы."),
    ("ТАНЦПОЛ", 100, "Как называется знаменитое движение,\nкогда кажется, что танцор скользит назад?", "Лунная походка", "Лунная походка! Можно показать для двойных аплодисментов."),
    ("ТАНЦПОЛ", 200, "Назовите песню, под которую ваша команда\nготова выйти на танцпол прямо сейчас.", "Творческий раунд", "Подходит любая песня — команда объясняет выбор за 10 секунд."),
    ("ТАНЦПОЛ", 300, "Покажите без слов танец человека,\nкоторый нашёл идеальный подарок.", "Пантомима", "Остальные команды угадывают. Самое точное или смешное исполнение получает баллы."),
    ("КИНО ВЕЧЕРА", 100, "Что обычно делают в зале,\nкогда начинается киносеанс?", "Выключают свет", "Выключают свет — и внимание на экран."),
    ("КИНО ВЕЧЕРА", 200, "Продолжите крылатую фразу:\n«Я вернусь...»", "«...и продолжу банкет»", "Засчитывается любой бодрый вариант. Классика — «Я вернусь!»"),
    ("КИНО ВЕЧЕРА", 300, "Придумайте название фильма\nо сегодняшнем дне рождения.", "Творческий раунд", "Самое афишное название выбирают именинники."),
    ("НА ВКУС", 100, "Какой ингредиент точно не помешает\nидеальному праздничному угощению?", "Настроение", "Верный ответ — настроение. А остальное можно заказать."),
    ("НА ВКУС", 200, "Придумайте название безалкогольного коктейля\nв честь Влада и Аллы.", "Творческий раунд", "За особенно яркое название — бонусные аплодисменты."),
    ("НА ВКУС", 300, "За 15 секунд назовите как можно больше\nначинок для пиццы.", "Блиц-раунд", "Каждая уникальная начинка — один балл. Повторы не считаются."),
]


def set_background(slide, color=NAVY):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def text_box(slide, text, x, y, w, h, size=24, color=WHITE, bold=False,
             align=PP_ALIGN.CENTER, font="Aptos Display"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def glow_bar(slide, x, y, w, color):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, Inches(.08))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_neon_decor(slide):
    for x, y, d, c in [
        (.35, .35, .18, CYAN), (12.72, .48, .12, PINK), (12.35, 6.9, .2, YELLOW),
        (.58, 6.82, .1, PINK), (11.75, .55, .08, YELLOW), (1.05, 6.95, .06, CYAN),
    ]:
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        dot.fill.solid()
        dot.fill.fore_color.rgb = c
        dot.line.fill.background()
    glow_bar(slide, Inches(.55), Inches(.48), Inches(2.0), CYAN)
    glow_bar(slide, Inches(10.78), Inches(6.98), Inches(1.95), PINK)


def card(slide, x, y, w, h, fill=DEEP_PURPLE, line=CYAN):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.8)
    return shape


def button(slide, label, x, y, w, h, color=CYAN, target=None):
    shape = card(slide, x, y, w, h, fill=NAVY, line=color)
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = "Aptos Display"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = color
    if target is not None:
        shape.click_action.target_slide = target
    return shape


def add_footer(slide, board, score):
    button(slide, "К ТАБЛО", Inches(.55), Inches(6.78), Inches(1.7), Inches(.42), CYAN, board)
    button(slide, "СЧЁТ", Inches(10.98), Inches(6.78), Inches(1.25), Inches(.42), PINK, score)


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, DEEP_PURPLE)
    add_neon_decor(slide)
    text_box(slide, "СВОЯ ИГРА", Inches(.7), Inches(1.15), Inches(11.9), Inches(.8), 40, CYAN, True)
    text_box(slide, "День рождения Влада и Аллы", Inches(.7), Inches(2.0), Inches(11.9), Inches(.55), 27, WHITE, True)
    text_box(slide, "Яркий вечер, громкие ответы и много поводов улыбнуться", Inches(1.2), Inches(2.73), Inches(10.9), Inches(.48), 18, MUTED)
    card(slide, Inches(3.2), Inches(4.05), Inches(6.93), Inches(1.12), fill=NAVY, line=PINK)
    text_box(slide, "10 команд  •  4 категории  •  12 вопросов", Inches(3.4), Inches(4.26), Inches(6.55), Inches(.35), 19, YELLOW, True)
    text_box(slide, "Нажмите «Начать»", Inches(4.5), Inches(5.75), Inches(4.35), Inches(.35), 16, MUTED)
    return slide


def add_rules_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_neon_decor(slide)
    text_box(slide, "КАК ИГРАЕМ", Inches(.7), Inches(.7), Inches(11.9), Inches(.55), 30, PINK, True)
    rules = [
        "Выберите вопрос на табло — стоимость равна числу баллов.",
        "На обсуждение даётся до 30 секунд.",
        "В творческих раундах решение — за именинниками или ведущим.",
        "Ведите счёт на отдельном слайде или в любом удобном месте.",
    ]
    for i, rule in enumerate(rules):
        y = Inches(1.55 + i * 1.05)
        card(slide, Inches(1.2), y, Inches(10.93), Inches(.72), fill=DEEP_PURPLE, line=CYAN if i % 2 == 0 else PINK)
        text_box(slide, f"{i + 1}.  {rule}", Inches(1.52), y + Inches(.1), Inches(10.25), Inches(.5), 18, WHITE, False, PP_ALIGN.LEFT)
    return slide


def add_board_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_neon_decor(slide)
    text_box(slide, "ТАБЛО", Inches(.65), Inches(.35), Inches(12.0), Inches(.5), 28, CYAN, True)
    return slide


def add_score_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_neon_decor(slide)
    text_box(slide, "СЧЁТ КОМАНД", Inches(.65), Inches(.43), Inches(12.0), Inches(.5), 28, PINK, True)
    for i in range(10):
        col, row = i % 2, i // 2
        x, y = Inches(.95 + col * 6.05), Inches(1.22 + row * 1.05)
        card(slide, x, y, Inches(5.3), Inches(.7), fill=DEEP_PURPLE, line=CYAN if col == 0 else PINK)
        text_box(slide, f"КОМАНДА {i + 1}", x + Inches(.18), y + Inches(.12), Inches(3.3), Inches(.4), 16, WHITE, True, PP_ALIGN.LEFT)
        text_box(slide, "0", x + Inches(4.1), y + Inches(.1), Inches(.85), Inches(.42), 21, YELLOW, True)
    text_box(slide, "Чтобы изменить названия и баллы: кликните по тексту прямо на этом слайде.", Inches(.95), Inches(6.63), Inches(11.3), Inches(.35), 13, MUTED)
    return slide


def add_question_slide(prs, category, points, question):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_neon_decor(slide)
    text_box(slide, category, Inches(.72), Inches(.58), Inches(8.5), Inches(.42), 18, PINK, True, PP_ALIGN.LEFT)
    text_box(slide, str(points), Inches(10.7), Inches(.44), Inches(1.75), Inches(.62), 28, YELLOW, True)
    card(slide, Inches(1.05), Inches(1.55), Inches(11.22), Inches(3.65), fill=DEEP_PURPLE, line=CYAN)
    text_box(slide, question, Inches(1.55), Inches(2.03), Inches(10.23), Inches(2.6), 28, WHITE, True)
    text_box(slide, "Время пошло!", Inches(4.73), Inches(5.74), Inches(3.85), Inches(.34), 16, MUTED, True)
    return slide


def add_answer_slide(prs, category, points, answer, note):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, DEEP_PURPLE)
    add_neon_decor(slide)
    text_box(slide, f"{category}  •  {points} БАЛЛОВ", Inches(.72), Inches(.58), Inches(11.9), Inches(.42), 17, CYAN, True, PP_ALIGN.LEFT)
    text_box(slide, "ОТВЕТ", Inches(.9), Inches(1.35), Inches(11.55), Inches(.53), 25, PINK, True)
    card(slide, Inches(1.25), Inches(2.12), Inches(10.83), Inches(1.36), fill=NAVY, line=YELLOW)
    text_box(slide, answer, Inches(1.63), Inches(2.38), Inches(10.1), Inches(.75), 29, YELLOW, True)
    text_box(slide, note, Inches(1.35), Inches(4.18), Inches(10.63), Inches(.95), 17, WHITE)
    return slide


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    title = add_title_slide(prs)
    rules = add_rules_slide(prs)
    board = add_board_slide(prs)
    score = add_score_slide(prs)

    question_slides = []
    answer_slides = []
    for category, points, question, answer, note in QUESTIONS:
        question_slides.append(add_question_slide(prs, category, points, question))
        answer_slides.append(add_answer_slide(prs, category, points, answer, note))

    button(title, "НАЧАТЬ", Inches(4.66), Inches(6.22), Inches(4.0), Inches(.62), PINK, rules)
    button(rules, "К ТАБЛО", Inches(4.66), Inches(6.25), Inches(4.0), Inches(.62), CYAN, board)
    add_footer(score, board, score)

    categories = ["ИМЕНИННИКИ", "ТАНЦПОЛ", "КИНО ВЕЧЕРА", "НА ВКУС"]
    col_x = [Inches(.62), Inches(3.81), Inches(7.0), Inches(10.19)]
    for col, category in enumerate(categories):
        card(board, col_x[col], Inches(1.18), Inches(2.52), Inches(.82), fill=DEEP_PURPLE, line=PINK)
        text_box(board, category, col_x[col] + Inches(.05), Inches(1.32), Inches(2.42), Inches(.4), 14, WHITE, True)
        for row, points in enumerate([100, 200, 300]):
            y = Inches(2.25 + row * 1.22)
            index = col * 3 + row
            button(board, str(points), col_x[col], y, Inches(2.52), Inches(.82), YELLOW, question_slides[index])
    button(board, "СЧЁТ", Inches(5.15), Inches(6.38), Inches(3.03), Inches(.55), PINK, score)

    for q_slide, a_slide in zip(question_slides, answer_slides):
        button(q_slide, "ПОКАЗАТЬ ОТВЕТ", Inches(4.32), Inches(6.22), Inches(4.7), Inches(.52), PINK, a_slide)
        add_footer(q_slide, board, score)
        button(a_slide, "К ТАБЛО", Inches(4.62), Inches(5.83), Inches(4.08), Inches(.58), CYAN, board)
        add_footer(a_slide, board, score)

    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
